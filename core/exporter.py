import os
import re
from pathlib import Path
from typing import Dict, Any, List, Optional

import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, Preformatted
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Register Thai-capable font on Windows
THAI_FONT_NAME = 'Helvetica'
THAI_FONT_BOLD = 'Helvetica-Bold'

for font_path in ['C:/Windows/Fonts/tahoma.ttf', 'C:/Windows/Fonts/leelawad.ttf', 'C:/Windows/Fonts/arial.ttf']:
    if os.path.exists(font_path):
        try:
            pdfmetrics.registerFont(TTFont('ThaiFont', font_path))
            THAI_FONT_NAME = 'ThaiFont'
            bold_path = font_path.replace('.ttf', 'bd.ttf')
            if os.path.exists(bold_path):
                pdfmetrics.registerFont(TTFont('ThaiFont-Bold', bold_path))
                THAI_FONT_BOLD = 'ThaiFont-Bold'
            else:
                THAI_FONT_BOLD = 'ThaiFont'
            break
        except Exception:
            pass

class DocumentExporter:
    """
    Handles inserting verified answers, tables, and diagrams into documents (.docx and .pdf).
    Flowchart: 'ใส่คำตอบ' -> 'output file เป็น .Doc/.pdf'
    """

    @classmethod
    def export_docx(
        cls,
        original_file_path: str,
        solved_items: List[Dict[str, Any]],
        output_path: str
    ) -> str:
        """
        Creates an updated .docx document with answers, tables, and diagrams.
        """
        is_original_docx = Path(original_file_path).suffix.lower() == ".docx"

        if is_original_docx and os.path.exists(original_file_path):
            try:
                doc = docx.Document(original_file_path)
                cls._fill_docx_blanks(doc, solved_items)
            except Exception:
                doc = cls._create_fresh_docx(solved_items, Path(original_file_path).name)
        else:
            doc = cls._create_fresh_docx(solved_items, Path(original_file_path).name)

        os.makedirs(Path(output_path).parent, exist_ok=True)
        doc.save(output_path)
        return output_path

    @classmethod
    def _fill_docx_blanks(cls, doc: docx.Document, solved_items: List[Dict[str, Any]]):
        """Replaces blank placeholders directly in existing docx paragraphs."""
        for item in solved_items:
            answer = item.get("answer", "")
            if not answer:
                continue

            for p in doc.paragraphs:
                if not p.text.strip():
                    continue

                if re.search(r'(\.{3,}|_{2,}|\(\s*\)|\[\s*\])', p.text):
                    new_text = re.sub(r'(\.{3,}|_{2,}|\(\s*\)|\[\s*\])', f" [{answer}] ", p.text, count=1)
                    if new_text != p.text:
                        p.text = new_text
                        for run in p.runs:
                            if answer in run.text:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(16, 110, 190)
                        break

        # Append solution & answer key summary section
        cls._append_solution_section(doc, solved_items)

    @classmethod
    def _create_fresh_docx(cls, solved_items: List[Dict[str, Any]], original_filename: str) -> docx.Document:
        doc = docx.Document()

        title = doc.add_heading("Homework Destroyer - Solved Assignment (เฉลยใบงาน)", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub = doc.add_paragraph(f"Source Document: {original_filename} | Total Solved: {len(solved_items)}")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

        cls._append_solution_section(doc, solved_items)
        return doc

    @classmethod
    def _append_solution_section(cls, doc: docx.Document, solved_items: List[Dict[str, Any]]):
        doc.add_page_break()
        heading = doc.add_heading("เฉลยใบงานและคำตอบที่ผ่านการตรวจสอบ (Verified Answer Key)", level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for item in solved_items:
            p_q = doc.add_paragraph()
            q_num = item.get("question_number") or f"ข้อ {item.get('id')}"
            r_num = p_q.add_run(f"{q_num}: ")
            r_num.bold = True
            r_num.font.color.rgb = RGBColor(30, 64, 175) # Dark Blue

            p_q.add_run(f"{item.get('raw_question')}\n")

            # Answer
            p_ans = doc.add_paragraph()
            r_ans_label = p_ans.add_run("คำตอบ: ")
            r_ans_label.bold = True
            r_ans_val = p_ans.add_run(f"{item.get('answer')}\n")
            r_ans_val.bold = True
            r_ans_val.font.color.rgb = RGBColor(16, 128, 55) # Green

            # If table data exists (e.g. Question 3 route table)
            table_data = item.get("table_data")
            if table_data and isinstance(table_data, list) and len(table_data) > 0:
                try:
                    p_tbl_lbl = doc.add_paragraph()
                    p_tbl_lbl.add_run("ตารางแสดงผลลัพธ์:\n").bold = True
                    
                    first_row = table_data[0]
                    if isinstance(first_row, dict):
                        headers = list(first_row.keys())
                        if headers:
                            t = doc.add_table(rows=len(table_data) + 1, cols=len(headers))
                            t.alignment = WD_TABLE_ALIGNMENT.CENTER
                            for c_idx, h in enumerate(headers):
                                cell = t.cell(0, c_idx)
                                cell.text = str(h)
                                for r in cell.paragraphs[0].runs:
                                    r.bold = True
                            for r_idx, row_dict in enumerate(table_data):
                                for c_idx, h in enumerate(headers):
                                    t.cell(r_idx + 1, c_idx).text = str(row_dict.get(h, ""))
                    elif isinstance(first_row, list):
                        num_cols = max(len(r) for r in table_data if isinstance(r, list))
                        if num_cols > 0:
                            t = doc.add_table(rows=len(table_data), cols=num_cols)
                            t.alignment = WD_TABLE_ALIGNMENT.CENTER
                            for r_idx, row_list in enumerate(table_data):
                                for c_idx, val in enumerate(row_list):
                                    t.cell(r_idx, c_idx).text = str(val)
                except Exception as te:
                    doc.add_paragraph(f"[Table: {str(table_data)}]")

            # If diagram exists
            diagram = item.get("diagram_ascii")
            if diagram:
                try:
                    p_dia = doc.add_paragraph()
                    p_dia.add_run("แผนผัง/โครงสร้างที่วาด:\n").bold = True
                    p_dia_box = doc.add_paragraph(diagram)
                    p_dia_box.paragraph_format.left_indent = Inches(0.5)
                except Exception:
                    doc.add_paragraph(str(diagram))

            # Explanation
            p_exp = doc.add_paragraph()
            p_exp.add_run("คำอธิบาย/วิธีคิด: ").bold = True
            p_exp.add_run(f"{item.get('explanation')}\n")

            # Verification badge
            status = "Verified (ตรวจสอบแล้ว ถูกต้อง)" if item.get("verified") else "Checked"
            rounds = item.get("verification_rounds", 0)
            p_meta = doc.add_paragraph()
            r_meta = p_meta.add_run(f"สถานะ: {status} | จำนวนรอบแก้ไข: {rounds}/2 | ความมั่นใจ: {item.get('confidence', 95)}%")
            r_meta.font.italic = True
            r_meta.font.color.rgb = RGBColor(120, 120, 120)

    @classmethod
    def export_pdf(
        cls,
        original_filename: str,
        solved_items: List[Dict[str, Any]],
        output_path: str
    ) -> str:
        """
        Creates a PDF document with Thai font support, tables, diagrams, and verification notes.
        """
        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName=THAI_FONT_BOLD,
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0f172a"),
            alignment=1,
            spaceAfter=6
        )

        sub_style = ParagraphStyle(
            'DocSub',
            parent=styles['Normal'],
            fontName=THAI_FONT_NAME,
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#64748b"),
            alignment=1,
            spaceAfter=14
        )

        q_head_style = ParagraphStyle(
            'QHead',
            parent=styles['Heading2'],
            fontName=THAI_FONT_BOLD,
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#1e40af"),
            spaceBefore=10,
            spaceAfter=4
        )

        body_style = ParagraphStyle(
            'QuestionBody',
            parent=styles['Normal'],
            fontName=THAI_FONT_NAME,
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6
        )

        answer_style = ParagraphStyle(
            'AnswerText',
            parent=styles['Normal'],
            fontName=THAI_FONT_BOLD,
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#15803d"),
            spaceAfter=4
        )

        meta_style = ParagraphStyle(
            'MetaText',
            parent=styles['Normal'],
            fontName=THAI_FONT_NAME,
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor("#64748b"),
            spaceAfter=8
        )

        elements = []

        elements.append(Paragraph("Homework Destroyer - Solved Assignment (เฉลยใบงาน)", title_style))
        elements.append(Paragraph(f"เอกสาร: {original_filename} | ทั้งหมด {len(solved_items)} ข้อ | Gemini AI Vision", sub_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#3b82f6"), spaceAfter=14))

        for item in solved_items:
            q_num = item.get("question_number") or f"ข้อ {item.get('id')}"
            raw = (item.get("raw_question") or "").replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            ans = (item.get("answer") or "").replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            exp = (item.get("explanation") or "").replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

            elements.append(Paragraph(f"<b>{q_num}:</b> {raw}", q_head_style))
            elements.append(Paragraph(f"<b>คำตอบ:</b> <font color='#15803d'>{ans}</font>", answer_style))

            # Table if any
            table_data = item.get("table_data")
            if table_data and isinstance(table_data, list) and len(table_data) > 0:
                try:
                    first_row = table_data[0]
                    if isinstance(first_row, dict):
                        headers = list(first_row.keys())
                        if headers:
                            table_matrix = [[Paragraph(f"<b>{str(h).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')}</b>", body_style) for h in headers]]
                            for row_dict in table_data:
                                table_matrix.append([Paragraph(str(row_dict.get(h, "")).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'), body_style) for h in headers])

                            report_tbl = Table(table_matrix, colWidths=[520 / len(headers)] * len(headers))
                            report_tbl.setStyle(TableStyle([
                                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#94a3b8")),
                                ('PADDING', (0, 0), (-1, -1), 4),
                                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                            ]))
                            elements.append(report_tbl)
                            elements.append(Spacer(1, 6))
                    elif isinstance(first_row, list):
                        num_cols = max(len(r) for r in table_data if isinstance(r, list))
                        if num_cols > 0:
                            table_matrix = []
                            for row_list in table_data:
                                table_matrix.append([Paragraph(str(v).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'), body_style) for v in row_list])
                            report_tbl = Table(table_matrix, colWidths=[520 / num_cols] * num_cols)
                            report_tbl.setStyle(TableStyle([
                                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#94a3b8")),
                                ('PADDING', (0, 0), (-1, -1), 4),
                            ]))
                            elements.append(report_tbl)
                            elements.append(Spacer(1, 6))
                except Exception as pte:
                    logger.warning(f"Error rendering PDF table: {pte}")

            # Diagram if any
            diagram = item.get("diagram_ascii")
            if diagram:
                try:
                    clean_diag = str(diagram).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    elements.append(Paragraph(f"<b>แผนผัง:</b><br/>{clean_diag}", body_style))
                    elements.append(Spacer(1, 4))
                except Exception:
                    pass

            elements.append(Paragraph(f"<b>วิธีคิด:</b> {exp}", body_style))

            status = "Verified [ผ่านการตรวจสอบ]" if item.get("verified") else "Checked"
            rounds = item.get("verification_rounds", 0)
            elements.append(Paragraph(f"สถานะ: {status} | รอบการแก้ไข: {rounds}/2 | ความมั่นใจ: {item.get('confidence', 95)}%", meta_style))
            elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e2e8f0"), spaceAfter=10))

        os.makedirs(Path(output_path).parent, exist_ok=True)
        doc.build(elements)
        return output_path
